"""Testes unitários para o parser ABDC."""

from pathlib import Path

import openpyxl
import pytest

from pipeline.parsers.base import ParsedRecord, SourceParser
from pipeline.parsers.parser_abdc import Parser


@pytest.fixture
def abdc_xlsx(tmp_path: Path) -> Path:
    """Cria um XLSX sintético no formato ABDC Journal Quality List."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2025 JQL"

    # Linhas de cabeçalho decorativo (simulando a planilha real)
    for _ in range(5):
        ws.append([None] * 8)
    ws.append([None, "2025 ABDC Journal Quality List"] + [None] * 6)
    ws.append(
        [None, "Description text"] + [None] * 6
    )

    # Cabeçalho real (row 8 na planilha, 0-indexed row 7)
    ws.append(
        [
            None,
            "Journal Title",
            "Publisher",
            "ISSN",
            "ISSNOnline",
            "Year Inception",
            "FoR",
            "2025 rating",
        ]
    )

    # Dados de teste
    ws.append([None, "Academy of Management Journal", "Academy of Management", "0001-4273", "1948-0989", "1958", 3507, "A*"])
    ws.append([None, "Abacus", "Wiley-Blackwell", "0001-3072", "1467-6281", "1965", 3501, "A"])
    ws.append([None, "Business Ethics Quarterly", "Cambridge UP", "1052-150X", "2153-3326", "1991", 3507, "B"])
    ws.append([None, "Asian Business & Management", "Springer", "1472-4782", "1476-9328", "2002", 3507, "C"])

    xlsx_path = tmp_path / "abdc_test.xlsx"
    wb.save(xlsx_path)
    wb.close()
    return xlsx_path


@pytest.fixture
def abdc_xlsx_with_whitespace(tmp_path: Path) -> Path:
    """XLSX com valores contendo espaços e tabs à direita (como na fonte real)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2025 JQL"

    for _ in range(5):
        ws.append([None] * 8)
    ws.append([None, "2025 ABDC Journal Quality List"] + [None] * 6)
    ws.append([None, "Desc"] + [None] * 6)
    ws.append(
        [None, "Journal Title", "Publisher", "ISSN", "ISSNOnline",
         "Year Inception", "FoR", "2025 rating"]
    )

    # Valores com trailing whitespace/tabs
    ws.append([None, "4OR ", "Springer Nature", "1619-4500", "1614-2411\t", "2003\t", 3599, "B "])
    ws.append([None, "Some Journal", "Publisher", "1111-2222\t", "\t", "2000\t", 3501, "A "])

    xlsx_path = tmp_path / "abdc_whitespace.xlsx"
    wb.save(xlsx_path)
    wb.close()
    return xlsx_path


@pytest.fixture
def abdc_xlsx_empty_rows(tmp_path: Path) -> Path:
    """XLSX com linhas vazias e ratings inválidos intercalados."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2025 JQL"

    for _ in range(5):
        ws.append([None] * 8)
    ws.append([None, "2025 ABDC Journal Quality List"] + [None] * 6)
    ws.append([None, "Desc"] + [None] * 6)
    ws.append(
        [None, "Journal Title", "Publisher", "ISSN", "ISSNOnline",
         "Year Inception", "FoR", "2025 rating"]
    )

    # Linha válida
    ws.append([None, "Valid Journal", "Pub", "1234-5678", "8765-4321", "2000", 3501, "A"])
    # Linha vazia (sem rating)
    ws.append([None, None, None, None, None, None, None, None])
    # Linha com rating inválido
    ws.append([None, "Bad Rating", "Pub", "9999-0000", None, "2010", 3501, "X"])
    # Outra linha válida
    ws.append([None, "Another Valid", "Pub", "5555-6666", None, "2015", 3501, "C"])

    xlsx_path = tmp_path / "abdc_empty.xlsx"
    wb.save(xlsx_path)
    wb.close()
    return xlsx_path


class TestParserInterface:
    """Verifica que o parser cumpre a interface SourceParser."""

    def test_source_id(self):
        parser = Parser()
        assert parser.source_id == "abdc"

    def test_provides(self):
        parser = Parser()
        assert parser.provides == ["abdc_rating"]

    def test_implements_protocol(self):
        """Verifica conformidade com SourceParser Protocol."""
        parser = Parser()
        assert hasattr(parser, "source_id")
        assert hasattr(parser, "provides")
        assert hasattr(parser, "parse")


class TestIssnExtraction:
    """Verifica extração de ISSNs."""

    def test_both_issns_extracted(self, abdc_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx)))
        rec = records[0]
        assert rec.issn == "0001-4273"
        assert rec.issn_alt == "1948-0989"

    def test_all_records_have_issn(self, abdc_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx)))
        for rec in records:
            assert rec.issn is not None


class TestRatingExtraction:
    """Verifica extração do rating ABDC."""

    def test_a_star_rating(self, abdc_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx)))
        assert records[0].metricas["abdc_rating"] == "A*"

    def test_a_rating(self, abdc_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx)))
        assert records[1].metricas["abdc_rating"] == "A"

    def test_b_rating(self, abdc_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx)))
        assert records[2].metricas["abdc_rating"] == "B"

    def test_c_rating(self, abdc_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx)))
        assert records[3].metricas["abdc_rating"] == "C"


class TestWhitespaceHandling:
    """Verifica limpeza de espaços e tabs à direita."""

    def test_issn_stripped(self, abdc_xlsx_with_whitespace: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx_with_whitespace)))
        assert records[0].issn == "1619-4500"
        assert records[0].issn_alt == "1614-2411"

    def test_rating_stripped(self, abdc_xlsx_with_whitespace: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx_with_whitespace)))
        assert records[0].metricas["abdc_rating"] == "B"
        assert records[1].metricas["abdc_rating"] == "A"

    def test_title_stripped(self, abdc_xlsx_with_whitespace: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx_with_whitespace)))
        assert records[0].titulo == "4OR"

    def test_empty_issn_online_is_none(self, abdc_xlsx_with_whitespace: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx_with_whitespace)))
        # Second record has ISSNOnline as just "\t" → should be None
        assert records[1].issn_alt is None


class TestEmptyAndInvalidRows:
    """Verifica que linhas vazias e ratings inválidos são ignorados."""

    def test_only_valid_records_yielded(self, abdc_xlsx_empty_rows: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx_empty_rows)))
        assert len(records) == 2

    def test_valid_records_correct(self, abdc_xlsx_empty_rows: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx_empty_rows)))
        assert records[0].issn == "1234-5678"
        assert records[0].metricas["abdc_rating"] == "A"
        assert records[1].issn == "5555-6666"
        assert records[1].metricas["abdc_rating"] == "C"

    def test_no_issn_alt_when_missing(self, abdc_xlsx_empty_rows: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx_empty_rows)))
        assert records[1].issn_alt is None


class TestParsedRecordType:
    """Verifica que o parser produz instâncias de ParsedRecord."""

    def test_yields_parsed_record_instances(self, abdc_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx)))
        assert len(records) == 4
        for rec in records:
            assert isinstance(rec, ParsedRecord)

    def test_titulo_extracted(self, abdc_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx)))
        assert records[0].titulo == "Academy of Management Journal"
        assert records[1].titulo == "Abacus"

    def test_doi_is_none(self, abdc_xlsx: Path):
        """ABDC não fornece DOI."""
        parser = Parser()
        records = list(parser.parse(str(abdc_xlsx)))
        for rec in records:
            assert rec.doi is None
