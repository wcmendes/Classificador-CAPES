"""Testes unitários para o parser Qualis Histórico (2021-2024)."""

from pathlib import Path

import openpyxl
import pytest

from pipeline.parsers.base import ParsedRecord, SourceParser
from pipeline.parsers.parser_qualis import Parser


@pytest.fixture
def qualis_xlsx(tmp_path: Path) -> Path:
    """Cria um XLSX sintético no formato Qualis publicado pela CAPES."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RelatorioQualis"

    # Cabeçalho
    ws.append(["ISSN", "Título", "Área de Avaliação", "Estrato"])

    # Dados de teste — diferentes áreas e estratos
    ws.append([
        "0001-0782",
        "Communications of the ACM",
        "CIÊNCIA DA COMPUTAÇÃO",
        "A1",
    ])
    ws.append([
        "1413-6090",
        "A Economia em Revista",
        "ADMINISTRAÇÃO PÚBLICA E DE EMPRESAS, CIÊNCIAS CONTÁBEIS E TURISMO",
        "B3",
    ])
    ws.append([
        "2236-2029",
        "A Economia em Revista - AERE",
        "ADMINISTRAÇÃO PÚBLICA E DE EMPRESAS, CIÊNCIAS CONTÁBEIS E TURISMO",
        "B3",
    ])
    ws.append([
        "1234-5678",
        "Journal of Testing",
        "CIÊNCIA DA COMPUTAÇÃO",
        "A4",
    ])
    ws.append([
        "9999-0000",
        "Periódico de Educação",
        "EDUCAÇÃO",
        "B1",
    ])

    xlsx_path = tmp_path / "qualis_test.xlsx"
    wb.save(xlsx_path)
    wb.close()
    return xlsx_path


@pytest.fixture
def qualis_xlsx_all_estratos(tmp_path: Path) -> Path:
    """XLSX com todos os estratos válidos para verificar cobertura."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RelatorioQualis"

    ws.append(["ISSN", "Título", "Área de Avaliação", "Estrato"])

    estratos = ["A1", "A2", "A3", "A4", "B1", "B2", "B3", "B4", "C"]
    for i, estrato in enumerate(estratos):
        ws.append([
            f"{i:04d}-0000",
            f"Journal {estrato}",
            "ÁREA TESTE",
            estrato,
        ])

    xlsx_path = tmp_path / "qualis_all_estratos.xlsx"
    wb.save(xlsx_path)
    wb.close()
    return xlsx_path


@pytest.fixture
def qualis_xlsx_invalid_rows(tmp_path: Path) -> Path:
    """XLSX com linhas inválidas: sem ISSN, sem estrato, estrato desconhecido."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RelatorioQualis"

    ws.append(["ISSN", "Título", "Área de Avaliação", "Estrato"])

    # Linha válida
    ws.append(["1111-2222", "Valid Journal", "ÁREA X", "A1"])
    # Linha sem ISSN
    ws.append([None, "No ISSN Journal", "ÁREA X", "A2"])
    # Linha sem estrato
    ws.append(["3333-4444", "No Estrato", "ÁREA X", None])
    # Linha com estrato inválido
    ws.append(["5555-6666", "Bad Estrato", "ÁREA X", "Z9"])
    # Linha com ISSN vazio (string vazia)
    ws.append(["", "Empty ISSN", "ÁREA X", "B1"])
    # Outra linha válida
    ws.append(["7777-8888", "Another Valid", "ÁREA Y", "C"])

    xlsx_path = tmp_path / "qualis_invalid.xlsx"
    wb.save(xlsx_path)
    wb.close()
    return xlsx_path


@pytest.fixture
def qualis_xlsx_whitespace(tmp_path: Path) -> Path:
    """XLSX com valores contendo espaços à direita."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RelatorioQualis"

    ws.append(["ISSN", "Título", "Área de Avaliação", "Estrato"])

    ws.append(["1111-2222 ", " Journal with Spaces ", " ÁREA X ", " A1 "])
    ws.append(["3333-4444", "Clean Journal", "ÁREA Y", "B2"])

    xlsx_path = tmp_path / "qualis_whitespace.xlsx"
    wb.save(xlsx_path)
    wb.close()
    return xlsx_path


class TestParserInterface:
    """Verifica que o parser cumpre a interface SourceParser."""

    def test_source_id(self):
        parser = Parser()
        assert parser.source_id == "qualis"

    def test_provides(self):
        parser = Parser()
        assert parser.provides == ["qualis_estrato_historico"]

    def test_implements_protocol(self):
        """Verifica conformidade com SourceParser Protocol."""
        parser = Parser()
        assert hasattr(parser, "source_id")
        assert hasattr(parser, "provides")
        assert hasattr(parser, "parse")


class TestBasicParsing:
    """Verifica extração básica de dados do XLSX."""

    def test_yields_correct_number_of_records(self, qualis_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx)))
        assert len(records) == 5

    def test_yields_parsed_record_instances(self, qualis_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx)))
        for rec in records:
            assert isinstance(rec, ParsedRecord)

    def test_issn_extracted(self, qualis_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx)))
        assert records[0].issn == "0001-0782"
        assert records[1].issn == "1413-6090"

    def test_titulo_extracted(self, qualis_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx)))
        assert records[0].titulo == "Communications of the ACM"
        assert records[1].titulo == "A Economia em Revista"

    def test_estrato_in_metricas(self, qualis_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx)))
        assert records[0].metricas["qualis_estrato_historico"] == "A1"
        assert records[1].metricas["qualis_estrato_historico"] == "B3"

    def test_area_avaliacao_in_metricas(self, qualis_xlsx: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx)))
        assert records[0].metricas["qualis_area_avaliacao"] == "CIÊNCIA DA COMPUTAÇÃO"
        assert records[1].metricas["qualis_area_avaliacao"] == (
            "ADMINISTRAÇÃO PÚBLICA E DE EMPRESAS, CIÊNCIAS CONTÁBEIS E TURISMO"
        )

    def test_issn_alt_is_none(self, qualis_xlsx: Path):
        """Qualis não fornece ISSN alternativo."""
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx)))
        for rec in records:
            assert rec.issn_alt is None

    def test_doi_is_none(self, qualis_xlsx: Path):
        """Qualis não fornece DOI."""
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx)))
        for rec in records:
            assert rec.doi is None


class TestAllEstratos:
    """Verifica que todos os estratos válidos são reconhecidos."""

    def test_all_valid_estratos_parsed(self, qualis_xlsx_all_estratos: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx_all_estratos)))
        assert len(records) == 9
        estratos = [r.metricas["qualis_estrato_historico"] for r in records]
        assert set(estratos) == {"A1", "A2", "A3", "A4", "B1", "B2", "B3", "B4", "C"}


class TestInvalidRows:
    """Verifica que linhas inválidas são corretamente ignoradas."""

    def test_only_valid_records_yielded(self, qualis_xlsx_invalid_rows: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx_invalid_rows)))
        # Apenas 2 linhas válidas: "1111-2222" com A1 e "7777-8888" com C
        assert len(records) == 2

    def test_valid_records_correct(self, qualis_xlsx_invalid_rows: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx_invalid_rows)))
        assert records[0].issn == "1111-2222"
        assert records[0].metricas["qualis_estrato_historico"] == "A1"
        assert records[1].issn == "7777-8888"
        assert records[1].metricas["qualis_estrato_historico"] == "C"


class TestWhitespaceHandling:
    """Verifica limpeza de espaços nos valores."""

    def test_issn_stripped(self, qualis_xlsx_whitespace: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx_whitespace)))
        assert records[0].issn == "1111-2222"

    def test_titulo_stripped(self, qualis_xlsx_whitespace: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx_whitespace)))
        assert records[0].titulo == "Journal with Spaces"

    def test_estrato_stripped(self, qualis_xlsx_whitespace: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx_whitespace)))
        assert records[0].metricas["qualis_estrato_historico"] == "A1"

    def test_area_stripped(self, qualis_xlsx_whitespace: Path):
        parser = Parser()
        records = list(parser.parse(str(qualis_xlsx_whitespace)))
        assert records[0].metricas["qualis_area_avaliacao"] == "ÁREA X"


class TestMultipleAreas:
    """Verifica que o parser preserva a área de avaliação por registro."""

    def test_same_issn_different_areas_are_separate_records(self, tmp_path: Path):
        """Um mesmo ISSN pode aparecer em múltiplas áreas no Qualis."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "RelatorioQualis"
        ws.append(["ISSN", "Título", "Área de Avaliação", "Estrato"])
        ws.append(["0001-0782", "Communications of the ACM", "CIÊNCIA DA COMPUTAÇÃO", "A1"])
        ws.append(["0001-0782", "Communications of the ACM", "ENGENHARIAS IV", "A2"])

        xlsx_path = tmp_path / "qualis_multi_area.xlsx"
        wb.save(xlsx_path)
        wb.close()

        parser = Parser()
        records = list(parser.parse(str(xlsx_path)))
        assert len(records) == 2
        assert records[0].metricas["qualis_area_avaliacao"] == "CIÊNCIA DA COMPUTAÇÃO"
        assert records[0].metricas["qualis_estrato_historico"] == "A1"
        assert records[1].metricas["qualis_area_avaliacao"] == "ENGENHARIAS IV"
        assert records[1].metricas["qualis_estrato_historico"] == "A2"
