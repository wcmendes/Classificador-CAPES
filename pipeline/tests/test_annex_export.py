"""Testes básicos para exportação de Anexos 3 e 4.

Valida a estrutura de colunas, cabeçalhos, indicadores de excedente,
preservação literal de títulos e campos de estrato.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from pipeline.annex_export import (
    AnnexItem,
    AnnexTechItem,
    generate_annex3,
    generate_annex4,
    _COLUMNS_ANNEX3,
    _COLUMNS_ANNEX4,
    _FOOTNOTE,
    _HEADER_ANNEX3,
    _SUBTITLE_ANNEX3,
    _SUBTITLE_ANNEX4,
)


@pytest.fixture
def sample_items() -> list[AnnexItem]:
    """Fixture com itens de exemplo para Anexo 3."""
    return [
        AnnexItem(
            tipo="periódico",
            issn_ou_isbn="1234-5678",
            sigla_evento="",
            titulo="A Novel Approach to Deep Learning",
            autores_docentes="Silva, J.",
            autores_discentes="Souza, M.",
            outros_autores="Wang, X.; Park, K.",
            justificativa="",
            estrato="A1",
        ),
        AnnexItem(
            tipo="evento",
            issn_ou_isbn="",
            sigla_evento="SBBD",
            titulo="Query Optimization in Distributed Databases",
            autores_docentes="Oliveira, A.",
            autores_discentes="",
            outros_autores="Chen, L.",
            justificativa="",
            estrato="A3",
        ),
        AnnexItem(
            tipo="livro",
            issn_ou_isbn="978-3-16-148410-0",
            sigla_evento="",
            titulo="Fundamentos de Computação Quântica",
            autores_docentes="Costa, R.",
            autores_discentes="Lima, P.",
            outros_autores="",
            justificativa="",
            estrato="NAO_CLASSIFICAVEL",
        ),
    ]


@pytest.fixture
def sample_tech_items() -> list[AnnexTechItem]:
    """Fixture com itens de exemplo para Anexo 4."""
    return [
        AnnexTechItem(
            tipo="Software/Aplicativo",
            titulo="Sistema de Análise Bibliométrica v2.0",
            autores_docentes="Silva, J.",
            autores_discentes="Souza, M.",
            outros_autores="",
            parceiro_institucional="UFMG",
            justificativa="",
            produto_bibliografico="Software/Aplicativo (Programa de computador)",
            estrato="A2",
        ),
    ]


class TestAnnex3Structure:
    """Testes de estrutura do Anexo 3."""

    def test_generates_xlsx_file(self, sample_items, tmp_path):
        """Arquivo XLSX é gerado no caminho especificado."""
        output = tmp_path / "anexo3.xlsx"
        result = generate_annex3(sample_items, n_docentes=10, output_path=output)
        assert result.exists()
        assert result.suffix == ".xlsx"

    def test_sheet_name(self, sample_items, tmp_path):
        """Nome da aba corresponde ao template oficial."""
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(sample_items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        assert "Anexo 3 - Prod Bibliográfica" in wb.sheetnames

    def test_header_row(self, sample_items, tmp_path):
        """Cabeçalho do MEC na primeira linha (célula B1)."""
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(sample_items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        assert ws.cell(row=1, column=2).value == _HEADER_ANNEX3

    def test_subtitle_row(self, sample_items, tmp_path):
        """Subtítulo do anexo na linha 3."""
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(sample_items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        assert ws.cell(row=3, column=2).value == _SUBTITLE_ANNEX3

    def test_column_headers(self, sample_items, tmp_path):
        """Colunas na linha 5 correspondem ao template."""
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(sample_items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        headers = [ws.cell(row=5, column=i).value for i in range(1, len(_COLUMNS_ANNEX3) + 1)]
        assert headers == _COLUMNS_ANNEX3

    def test_data_starts_at_row_6(self, sample_items, tmp_path):
        """Dados começam na linha 6."""
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(sample_items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        assert ws.cell(row=6, column=1).value == 1
        assert ws.cell(row=6, column=2).value == "periódico"

    def test_merged_cells(self, sample_items, tmp_path):
        """Células mescladas no cabeçalho conforme template."""
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(sample_items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        merged = [str(r) for r in ws.merged_cells.ranges]
        assert "B1:K1" in merged
        assert "B3:G3" in merged


class TestAnnex3TitlePreservation:
    """Títulos devem ser exportados byte-a-byte (Req 14.2)."""

    def test_title_preserved_literally(self, tmp_path):
        """Título com caracteres especiais é preservado sem alteração."""
        title = "Análise de REDES com símbolos: α, β, γ — e hífens"
        items = [
            AnnexItem(
                tipo="periódico",
                issn_ou_isbn="0000-0001",
                titulo=title,
                estrato="A1",
            )
        ]
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        assert ws.cell(row=6, column=5).value == title

    def test_title_no_capitalization_correction(self, tmp_path):
        """Título com capitalização irregular NÃO é corrigido."""
        title = "a STUDY on lowercase AND UPPER"
        items = [
            AnnexItem(tipo="periódico", issn_ou_isbn="0000-0002", titulo=title, estrato="A2")
        ]
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        assert ws.cell(row=6, column=5).value == title


class TestAnnex3ExceedsLimit:
    """Itens excedentes de 4N marcados com indicador visual (Req 14.3)."""

    def test_items_within_limit_not_marked(self, tmp_path):
        """Itens dentro do limite 4N não têm indicador."""
        items = [
            AnnexItem(tipo="periódico", issn_ou_isbn="0000-0001", titulo=f"Title {i}", estrato="A1")
            for i in range(4)
        ]
        output = tmp_path / "anexo3.xlsx"
        # n_docentes=1 → limit=4, all within limit
        generate_annex3(items, n_docentes=1, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        for row in range(6, 10):
            val = ws.cell(row=row, column=11).value
            assert val in (None, ""), f"Row {row} should not be marked, got {val!r}"

    def test_items_exceeding_limit_marked(self, tmp_path):
        """Itens além de 4N são marcados com 'SIM'."""
        items = [
            AnnexItem(tipo="periódico", issn_ou_isbn="0000-0001", titulo=f"Title {i}", estrato="A1")
            for i in range(6)
        ]
        output = tmp_path / "anexo3.xlsx"
        # n_docentes=1 → limit=4, items 5 and 6 exceed
        generate_annex3(items, n_docentes=1, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        # Items 1-4 within limit
        for row in range(6, 10):
            val = ws.cell(row=row, column=11).value
            assert val in (None, ""), f"Row {row} should not be marked, got {val!r}"
        # Items 5-6 exceed
        assert ws.cell(row=10, column=11).value == "SIM"
        assert ws.cell(row=11, column=11).value == "SIM"

    def test_exceeding_items_not_discarded(self, tmp_path):
        """Itens excedentes são mantidos no arquivo (não descartados)."""
        items = [
            AnnexItem(tipo="periódico", issn_ou_isbn="0000-0001", titulo=f"Title {i}", estrato="A1")
            for i in range(6)
        ]
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(items, n_docentes=1, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        # All 6 items present
        assert ws.cell(row=11, column=1).value == 6  # 6th item


class TestAnnex3EstratoField:
    """NAO_CLASSIFICAVEL e NAO_CONSIDERADO no campo de estrato (Req 14.5)."""

    def test_nao_classificavel_in_estrato(self, tmp_path):
        """Item NAO_CLASSIFICAVEL mostra status no campo estrato."""
        items = [
            AnnexItem(tipo="livro", issn_ou_isbn="978-0-00-000000-0", titulo="Book", estrato="NAO_CLASSIFICAVEL")
        ]
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        assert ws.cell(row=6, column=10).value == "NAO_CLASSIFICAVEL"

    def test_nao_considerado_in_estrato(self, tmp_path):
        """Item NAO_CONSIDERADO mostra status no campo estrato."""
        items = [
            AnnexItem(tipo="evento", sigla_evento="XYZ", titulo="Paper", estrato="NAO_CONSIDERADO")
        ]
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        assert ws.cell(row=6, column=10).value == "NAO_CONSIDERADO"


class TestAnnex3Footnote:
    """Nota de rodapé sobre autoria múltipla (Req 14.4)."""

    def test_footnote_present(self, sample_items, tmp_path):
        """Nota de rodapé aparece após os dados."""
        output = tmp_path / "anexo3.xlsx"
        generate_annex3(sample_items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        # Footnote row = len(items) + 5 + 2 = 3 + 5 + 2 = 10
        footnote_row = len(sample_items) + 5 + 2
        cell_value = ws.cell(row=footnote_row, column=1).value
        assert _FOOTNOTE in cell_value


class TestAnnex4Structure:
    """Testes de estrutura do Anexo 4."""

    def test_generates_xlsx_file(self, sample_tech_items, tmp_path):
        """Arquivo XLSX é gerado no caminho especificado."""
        output = tmp_path / "anexo4.xlsx"
        result = generate_annex4(sample_tech_items, n_docentes=10, output_path=output)
        assert result.exists()

    def test_sheet_name(self, sample_tech_items, tmp_path):
        """Nome da aba corresponde ao template oficial."""
        output = tmp_path / "anexo4.xlsx"
        generate_annex4(sample_tech_items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        assert "Anexo 4 - Prod Técnica" in wb.sheetnames

    def test_column_headers(self, sample_tech_items, tmp_path):
        """Colunas na linha 5 correspondem ao template."""
        output = tmp_path / "anexo4.xlsx"
        generate_annex4(sample_tech_items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 4 - Prod Técnica"]
        headers = [ws.cell(row=5, column=i).value for i in range(1, len(_COLUMNS_ANNEX4) + 1)]
        assert headers == _COLUMNS_ANNEX4

    def test_subtitle_row(self, sample_tech_items, tmp_path):
        """Subtítulo do Anexo 4 na linha 3."""
        output = tmp_path / "anexo4.xlsx"
        generate_annex4(sample_tech_items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 4 - Prod Técnica"]
        assert ws.cell(row=3, column=2).value == _SUBTITLE_ANNEX4

    def test_limit_m_calculation(self, tmp_path):
        """M = max(10, N/4). Com N=80, M=20."""
        items = [
            AnnexTechItem(tipo="Software", titulo=f"Soft {i}", estrato="A1")
            for i in range(25)
        ]
        output = tmp_path / "anexo4.xlsx"
        generate_annex4(items, n_docentes=80, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 4 - Prod Técnica"]
        # M = max(10, 80//4) = 20; items 21-25 exceed
        for row in range(6, 26):  # items 1-20 within limit
            val = ws.cell(row=row, column=10).value
            assert val in (None, ""), f"Row {row} should not be marked, got {val!r}"
        # Items 21-25 exceed
        assert ws.cell(row=26, column=10).value == "SIM"

    def test_limit_m_minimum_10(self, tmp_path):
        """M nunca é menor que 10, mesmo com poucos docentes."""
        items = [
            AnnexTechItem(tipo="Software", titulo=f"Soft {i}", estrato="A1")
            for i in range(12)
        ]
        output = tmp_path / "anexo4.xlsx"
        # n_docentes=8 → N/4=2, M=max(10,2)=10; items 11-12 exceed
        generate_annex4(items, n_docentes=8, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 4 - Prod Técnica"]
        val = ws.cell(row=15, column=10).value  # item 10 within limit
        assert val in (None, ""), f"Item 10 should not be marked, got {val!r}"
        assert ws.cell(row=16, column=10).value == "SIM"  # item 11 exceeds

    def test_footnote_present(self, sample_tech_items, tmp_path):
        """Nota de rodapé presente no Anexo 4."""
        output = tmp_path / "anexo4.xlsx"
        generate_annex4(sample_tech_items, n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 4 - Prod Técnica"]
        footnote_row = len(sample_tech_items) + 5 + 2
        cell_value = ws.cell(row=footnote_row, column=1).value
        assert _FOOTNOTE in cell_value


class TestAnnex3EmptyList:
    """Geração com lista vazia de itens."""

    def test_empty_items_generates_valid_file(self, tmp_path):
        """Lista vazia gera arquivo válido com cabeçalhos e rodapé."""
        output = tmp_path / "anexo3_empty.xlsx"
        generate_annex3([], n_docentes=10, output_path=output)
        wb = load_workbook(output)
        ws = wb["Anexo 3 - Prod Bibliográfica"]
        # Headers present
        assert ws.cell(row=5, column=1).value == "No."
        # Footnote present
        footnote_row = 0 + 5 + 2  # 7
        assert _FOOTNOTE in ws.cell(row=footnote_row, column=1).value
