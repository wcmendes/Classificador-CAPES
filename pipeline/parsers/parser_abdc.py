"""Parser para a fonte ABDC Journal Quality List.

Formato do arquivo: XLSX (openpyxl).
Fonte restrita — exige path configurado em config/local.yml.

A planilha "2025 JQL" tem cabeçalho na linha 8 (0-indexed: row 7) com colunas:
  Journal Title | Publisher | ISSN | ISSNOnline | Year Inception | FoR | 2025 rating

Ratings válidos: A*, A, B, C.
"""

from typing import Iterator

import openpyxl

from pipeline.parsers.base import ParsedRecord

# Linha (0-indexed) onde ficam os cabeçalhos na planilha original
_HEADER_ROW = 7
_SHEET_NAME = "2025 JQL"


class Parser:
    """Parser ABDC implementando a interface SourceParser.

    Extrai ISSN(s), título e rating (A*, A, B, C) de cada linha
    do XLSX disponibilizado pela ABDC.
    """

    @property
    def source_id(self) -> str:
        return "abdc"

    @property
    def provides(self) -> list[str]:
        return ["abdc_rating"]

    def parse(self, path: str) -> Iterator[ParsedRecord]:
        """Lê o XLSX da ABDC e produz ParsedRecord por linha.

        Valores podem conter espaços e tabs à direita que são removidos.
        Linhas sem rating válido são ignoradas.
        """
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

        # Tentar a sheet nomeada; se não existir, usar a primeira
        if _SHEET_NAME in wb.sheetnames:
            ws = wb[_SHEET_NAME]
        else:
            ws = wb[wb.sheetnames[0]]

        header_map: dict[str, int] | None = None

        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if header_map is None:
                # Procurar a linha de cabeçalho (contém "Journal Title")
                cells = [str(c).strip() if c is not None else "" for c in row]
                if "Journal Title" in cells:
                    header_map = {cells[i]: i for i in range(len(cells)) if cells[i]}
                continue

            # Extrair valores usando o mapa de cabeçalhos
            titulo = self._cell_str(row, header_map.get("Journal Title"))
            issn = self._cell_str(row, header_map.get("ISSN"))
            issn_online = self._cell_str(row, header_map.get("ISSNOnline"))
            rating = self._cell_str(row, header_map.get("2025 rating"))

            # Ignorar linhas vazias ou sem rating válido
            if not rating:
                continue

            # Normalizar rating (remover espaços)
            rating = rating.strip()
            if rating not in ("A*", "A", "B", "C"):
                continue

            # Montar métricas
            metricas: dict[str, str | int | float] = {"abdc_rating": rating}

            yield ParsedRecord(
                issn=issn or None,
                issn_alt=issn_online or None,
                doi=None,
                titulo=titulo or None,
                metricas=metricas,
            )

        wb.close()

    @staticmethod
    def _cell_str(row: tuple, col_idx: int | None) -> str | None:
        """Extrai valor de célula como string limpa, ou None se vazia."""
        if col_idx is None or col_idx >= len(row):
            return None
        val = row[col_idx]
        if val is None:
            return None
        cleaned = str(val).strip().strip("\t")
        return cleaned if cleaned else None
